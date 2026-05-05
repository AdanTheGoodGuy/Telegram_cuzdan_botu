import os
import sqlite3
import logging
import random
import urllib.request
import xml.etree.ElementTree as ET
import json
import tempfile
from datetime import datetime, timedelta
from dotenv import load_dotenv
from telegram import Update, InlineKeyboardButton, InlineKeyboardMarkup
from telegram.ext import (
    Application, CommandHandler, MessageHandler, filters, ContextTypes, 
    ConversationHandler, CallbackQueryHandler
)
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# Logging ayarları
logging.basicConfig(format="%(asctime)s - %(name)s - %(levelname)s - %(message)s", level=logging.INFO)

# Sohbet durumları
WAITING_TYPE, WAITING_AMOUNT, WAITING_CATEGORY, WAITING_DESC = range(4)
EDIT_CHOOSE, EDIT_AMOUNT, EDIT_CATEGORY, EDIT_DESC = range(4, 8)

# Token yükleme
load_dotenv()
TOKEN = os.getenv("TELEGRAM_BOT_TOKEN")

# Sabitler
KATEGORILER = {
    "🛒 Market": "Market",
    "🏠 Fatura": "Fatura",
    "🚗 Ulaşım": "Ulaşım",
    "🍔 Yemek": "Yemek",
    "🎉 Eğlence": "Eğlence",
    "💊 Sağlık": "Sağlık",
    "📦 Diğer": "Diğer"
}
KATEGORI_CLEAN = list(KATEGORILER.values())

# Yardımcı: Düzenleme menü klavyesi
def edit_menu_keyboard():
    return InlineKeyboardMarkup([
        [InlineKeyboardButton("💸 Tutar", callback_data="edit_tutar"), InlineKeyboardButton("📁 Kategori", callback_data="edit_kategori")],
        [InlineKeyboardButton("📝 Açıklama", callback_data="edit_aciklama"), InlineKeyboardButton("✅ Tamamla", callback_data="edit_bitir")]
    ])

# Yardımcı: Kategori butonları
def category_buttons(prefix="edit_cat_"):
    return [[InlineKeyboardButton(d, callback_data=f"{prefix}{c}")] for d, c in KATEGORILER.items()]

# 💬 BAĞLAMSAL MESAJLAR
GELIR_MESAJLARI = [
    "Harika! Kazancın arttıkça özgürlüğün de artar. 💪💰",
    "Gelir eklendi, cüzdanın derin nefes aldı! 🌬️✨",
    "Kazanmak güzel, akıllıca yönetmek daha güzel. 🧠📈",
    "Bu ayın kazançları umut verici! Gelecek parlak. 🚀",
    "Her kuruşun hesabını bilmek, servetin ilk adımı. 👣💎",
    "Kazandığın için değil, biriktirdiğin için zenginsin. 🏦",
    "Cebine giren para, hedeflerine giden yoldur. 🛣️🌟",
    "Para kazanmak bir sanat, harcamak bir bilim. 🎨🔬"
]

GIDER_MESAJLARI = [
    "Harcama kaydedildi. Bilinçli harcama, birikimin yarısıdır. 🧠💡",
    "Paranın nereye gittiğini bilmek, zenginliğin ilk adımı. 👣📊",
    "Bu harcama gerekli miydi? Bir dahakine biraz daha düşünelim. 🤔⚖️",
    "Küçük harcamalar, büyük birikimlerin sessiz düşmanıdır. 🐜📉",
    "Not edildi! Bütçeni aşmamaya dikkat et. 📝⚠️",
    "Her harcama bir tercih, her tercih bir gelecektir. 🔮💸",
    "Harcamanı kaydettin, şimdi telafisi için plan yap. 📋🔄",
    "Para akıyor, ama senin kontrolünde mi? 🌊🎛️"
]

POZITIF_MESAJLAR = [
    "Bakiyen pozitif! Finansal hedeflerine doğru sağlam adımlarla ilerliyorsun. 🎯🏔️",
    "Tebrikler, mali durumun gayet sağlıklı! Bu tempoyu koru. 🌿📈",
    "Bu ivmeyi korursan hayallerine çok yakınsın. Yıldızlara doğru! 🚀🌌",
    "Finansal özgürlük yolda! Harika gidiyorsun. 🛣️✨",
    "Cüzdanın şişkin, ruhun huzurlu. Böyle devam! 💼🕊️",
    "Birikimlerin seni geleceğe taşıyacak. Sabırlı ol. 🐢💎"
]

NEGATIF_MESAJLAR = [
    "Bakiye eksiye düştü. Biraz daha tutumlu olmanın zamanı. ⚠️📉",
    "Dikkat! Harcamalarını gözden geçirip bütçe yapabilirsin. 📝🛑",
    "Kısa vadeli sıkıntı, uzun vadeli ders olsun. 💪📚",
    "Kırmızı bölgedesin. Acil durum planı yapmalısın. 🚨🗺️",
    "Harcamaları kısma zamanı gelmiş olabilir. 📉🔒",
    "Bütçeni tekrar gözden geçirmen faydalı olacak. 🧐📊"
]

FINANS_IPUCLARI = [
    "💡 İpucu: '50/30/20 kuralı'nı dene! Gelirinin %50'si ihtiyaçlara, %30'u isteklere, %20'si birikime.",
    "💡 İpucu: Küçük harcamaları not etmek, ay sonunda büyük fark yaratır. 🐜➡️🐘",
    "💡 İpucu: Acil durum fonu için en az 3 aylık giderini biriktirmeyi hedefle. 🚨💰",
    "💡 İpucu: Faturalarını otomatik ödemeye bağlamak, gecikme cezasını engeller. ⚡🔄",
    "💡 İpucu: Harcamadan önce 24 saat bekle. İhtiyaç mı, istek mi ayırt edeceksin. ⏳🤔",
    "💡 İpucu: Aylık harcama limiti koymak, kontrolü ele almanın en kolay yoludur. 📏🛑",
    "💡 İpucu: Markete aç karnına gitme! Liste yap, sadık kal. 🛒📝",
    "💡 İpucu: Kredi kartı borcunu her ay tam öde, faiz tuzağına düşme. 💳⚠️",
    "💡 İpucu: Gelirinin en az %10'unu 'önce kendine öde' olarak biriktir. 🏦🐖",
    "💡 İpucu: Harcamalarını kategorize etmek, israfı görünür kılar. 📂👁️"
]

ARAYICI_MESAJLARI = [
    "İşte aradığın kategori raporu! 📂🔍",
    "Bakalım bu kategoride ne kadar harcamışsın? 🧐💸",
    "Harcamaların detayları karşında! 📊🔎",
    "Bu kategori seni zorluyor olabilir mi? 🤔📉"
]

DUZENLE_MESAJLARI = [
    "✅ İşlem başarıyla güncellendi! Kayıtların hep doğru kalsın. 📝✨",
    "🔄 Düzenleme tamamlandı. Finansal geçmişin artık daha net! 📊👌",
    "🖊️ Değişiklik kaydedildi. Kontrol sende! 🎛️💯",
    "📋 İşlem revize edildi. Her detay önemli! 🔍💎"
]

IPTAL_HATIRLATMA = "\n\n🔙 İptal için `/iptal` yaz."

# Veritabanı Yardımcıları
def get_db():
    conn = sqlite3.connect('data.db')
    conn.row_factory = sqlite3.Row
    return conn

def init_db():
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute('''
        CREATE TABLE IF NOT EXISTS transactions (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id INTEGER NOT NULL,
            type TEXT NOT NULL,
            amount REAL NOT NULL,
            aciklama TEXT
        )
    ''')
    try: cursor.execute('ALTER TABLE transactions ADD COLUMN category TEXT DEFAULT "Genel"')
    except: pass
    try: cursor.execute('ALTER TABLE transactions ADD COLUMN date TEXT DEFAULT ""')
    except: pass
    conn.commit()
    conn.close()

# --- Komut İşleyicileri ---

async def start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    text = (
        "Merhaba! 👋 *Kişisel Finans Botuna* hoş geldin.\n\n"
        "Bu bot senin *dijital cüzdanın* 📱. Gelir ve giderlerini kolayca takip et, "
        "nerelere harcadığını gör ve finansal hedeflerine bir adım daha yaklaş. 🎯💰\n\n"
        "━━━━━━━━━━━━━━━━━━━\n"
        "🔰 *YENİ BAŞLAYANLAR İÇİN:*\n"
        "1️⃣ `/ekle` ile ilk gelir/gider işlemini kaydet\n"
        "2️⃣ `/bakiye` ile toplam durumunu kontrol et\n"
        "3️⃣ `/listele` ile son işlemlerini gözden geçir\n\n"
        "━━━━━━━━━━━━━━━━━━━\n"
        "📋 *TÜM KOMUTLAR:*\n\n"
        "📊 *İşlem Yönetimi:*\n"
        "➕ `/ekle` - Gelir veya gider ekle (Adım adım rehber)\n"
        "📜 `/listele` - Son 5 işlemini görüntüle\n"
        "🖊️ `/duzenle` <id> - Bir işlemi düzenle (tutar, kategori, açıklama)\n"
        "🗑️ `/sil` <id> - Bir işlemi sil\n\n"
        "📈 *Analiz ve Raporlar:*\n"
        "💳 `/bakiye` - Toplam gelir, gider ve net bakiyeni gör\n"
        "📅 `/haftalik` - Son 7 günün harcama özeti ve ipucu\n"
        "🔍 `/ara` - Kategoriye göre harcama detaylarını görüntüle\n"
        "📊 `/aylik_ozet` - Aylık karşılaştırmalı Excel raporunu indir\n\n"
        "💱 *Ekstra:*\n"
        "🏛️ `/kur` - TCMB günlük döviz ve altın kurları\n"
        "🆘 `/yardim` - Bu menüyü tekrar getir\n\n"
        "━━━━━━━━━━━━━━━━━━━\n"
        "💡 *İpucu:* Her zaman `/iptal` yazarak bir işlemi yarıda bırakabilirsin.\n\n"
        "Hadi başlayalım! 🚀 `/ekle` yazarak ilk işlemini kaydet."
    )
    await update.message.reply_text(text, parse_mode="Markdown")

async def yardim(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await start(update, context)

# --- DÖVİZ KURU KOMUTU ---
async def kur(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text("⏳ TCMB kurları alınıyor, lütfen bekle... ⌛")
    
    try:
        url = "https://www.tcmb.gov.tr/kurlar/today.xml"
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req, timeout=10) as response:
            xml_data = response.read().decode('utf-8')
            
        root = ET.fromstring(xml_data)
        
        def get_rate(currency_code):
            for child in root:
                if child.get('CurrencyCode') == currency_code:
                    rate = child.findtext('BanknoteSelling') or child.findtext('ForexSelling')
                    if rate:
                        return float(rate)
            return None

        usd = get_rate('USD')
        eur = get_rate('EUR')
        gbp = get_rate('GBP')
        
        gram_altin_str = "🥇 Gram Altın: _Veri yok_"
        try:
            gold_url = "https://www.vang.today/api/prices?type=XAUUSD"
            gold_req = urllib.request.Request(gold_url, headers={'User-Agent': 'Mozilla/5.0'})
            with urllib.request.urlopen(gold_req, timeout=10) as gold_response:
                gold_data = json.loads(gold_response.read().decode())
                if gold_data.get('success') and gold_data.get('buy') > 0 and usd:
                    ons_altin_usd = gold_data['buy']
                    gram_altin = (ons_altin_usd * usd) / 31.1035
                    gram_altin_str = f"🥇 Gram Altın: `{gram_altin:.2f}` TL"
        except Exception:
            pass
            
        usd_str = f"🇺🇸 1 Dolar: `{usd:.2f}` TL" if usd else "🇺🇸 1 Dolar: _Veri yok_"
        eur_str = f"🇪🇺 1 Euro: `{eur:.2f}` TL" if eur else "🇪🇺 1 Euro: _Veri yok_"
        gbp_str = f"🇬🇧 1 Sterlin: `{gbp:.2f}` TL" if gbp else "🇬🇧 1 Sterlin: _Veri yok_"
        
        msg = (
            f"🏛️ *TCMB Günlük Kurları*\n\n"
            f"{usd_str}\n"
            f"{eur_str}\n"
            f"{gbp_str}\n"
            f"{gram_altin_str}\n\n"
            f"⚠️ _Döviz kurları TCMB, altın verileri uluslararası piyasadır._"
        )
        await update.message.reply_text(msg, parse_mode="Markdown")
            
    except Exception as e:
        logging.error(f"Kur hatası: {e}")
        await update.message.reply_text("❌ TCMB'ye bağlanılamadı, daha sonra tekrar dene.")

# --- /ekle Sohbet Akışı ---

async def ekle_basla(update: Update, context: ContextTypes.DEFAULT_TYPE):
    keyboard = [
        [InlineKeyboardButton("💰 Gelir", callback_data="gelir"), InlineKeyboardButton("💸 Gider", callback_data="gider")]
    ]
    await update.message.reply_text("💼 *İşlem türünü seç:*" + IPTAL_HATIRLATMA, reply_markup=InlineKeyboardMarkup(keyboard), parse_mode="Markdown")
    return WAITING_TYPE

async def tur_secildi(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    context.user_data['type'] = query.data
    icon = "💰" if query.data == "gelir" else "💸"
    await query.edit_message_text(f"{icon} Seçim: *{query.data.upper()}*. \n💵 Şimdi tutarı gir (sadece sayı):" + IPTAL_HATIRLATMA, parse_mode="Markdown")
    return WAITING_AMOUNT

async def tutar_alindi(update: Update, context: ContextTypes.DEFAULT_TYPE):
    try:
        val = float(update.message.text.replace(",", "."))
        if val <= 0:
            await update.message.reply_text("⚠️ Lütfen sıfırdan büyük bir sayı gir. 📏" + IPTAL_HATIRLATMA, parse_mode="Markdown")
            return WAITING_AMOUNT
        context.user_data['amount'] = val
        keyboard = category_buttons()
        await update.message.reply_text(f"✅ Tutar: `{val} TL` kaydedildi.\n📂 Şimdi kategori seç:" + IPTAL_HATIRLATMA, reply_markup=InlineKeyboardMarkup(keyboard), parse_mode="Markdown")
        return WAITING_CATEGORY
    except ValueError:
        await update.message.reply_text("❌ Geçersiz tutar. Lütfen sayı gir (örn: 50 veya 50.5). 🔢" + IPTAL_HATIRLATMA, parse_mode="Markdown")
        return WAITING_AMOUNT

async def kategori_secildi(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    context.user_data['category'] = query.data
    await query.edit_message_text(f"📂 Kategori: *{query.data}*. \n📝 Açıklama yaz (veya 'atla' yaz):" + IPTAL_HATIRLATMA, parse_mode="Markdown")
    return WAITING_DESC

async def aciklama_alindi(update: Update, context: ContextTypes.DEFAULT_TYPE):
    try:
        user_id = update.effective_user.id
        data = context.user_data
        if not data.get('type') or data.get('amount') is None:
            await update.message.reply_text("⚠️ Veri hatası oluştu, lütfen `/ekle` ile tekrar başla. 🔄")
            return ConversationHandler.END

        desc = update.message.text if update.message.text.lower() != "atla" else ""
        
        conn = get_db()
        cursor = conn.cursor()
        cursor.execute(
            "INSERT INTO transactions (user_id, type, amount, category, aciklama, date) VALUES (?, ?, ?, ?, ?, ?)",
            (user_id, data['type'], data['amount'], data['category'], desc, datetime.now().strftime("%Y-%m-%d %H:%M:%S"))
        )
        conn.commit()
        conn.close()
        
        icon = "💰" if data['type'] == 'gelir' else "💸"
        await update.message.reply_text(
            f"✅ *İşlem Eklendi:*\n{icon} {data['type'].upper()} - `{data['amount']}` TL\n📂 Kategori: {data['category']}",
            parse_mode="Markdown"
        )
        
        msg = random.choice(GELIR_MESAJLARI if data['type'] == 'gelir' else GIDER_MESAJLARI)
        await update.message.reply_text(msg)
        
    except Exception as e:
        logging.error(f"Kayıt hatası: {e}")
        await update.message.reply_text("❌ İşlem kaydedilirken bir hata oluştu. Lütfen tekrar dene. 🔄\n💡 Sorun devam ederse `/ekle` ile yeniden başla.")
    finally:
        context.user_data.clear()
        return ConversationHandler.END

async def iptal(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text("❌ İşlem iptal edildi. Başa dönüyorsun... 🔙")
    context.user_data.clear()
    await start(update, context)
    return ConversationHandler.END

# --- Diğer Komutlar ---

async def bakiye(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user_id = update.effective_user.id
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("SELECT SUM(amount) FROM transactions WHERE user_id=? AND type='gelir'", (user_id,))
    gelir = cursor.fetchone()[0] or 0
    cursor.execute("SELECT SUM(amount) FROM transactions WHERE user_id=? AND type='gider'", (user_id,))
    gider = cursor.fetchone()[0] or 0
    conn.close()
    
    bakiye = gelir - gider
    msg = f"💳 *Bakiye Durumu*\n\n💰 Toplam Gelir: `{gelir}` TL\n💸 Toplam Gider: `{gider}` TL\n📊 Net Bakiye: `{bakiye}` TL"
    
    tip = random.choice(POZITIF_MESAJLAR if bakiye >= 0 else NEGATIF_MESAJLAR)
    msg += f"\n\n💬 _{tip}_"
    
    await update.message.reply_text(msg, parse_mode="Markdown")

async def haftalik(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user_id = update.effective_user.id
    limit_date = (datetime.now() - timedelta(days=7)).strftime("%Y-%m-%d %H:%M:%S")
    
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("SELECT SUM(amount) FROM transactions WHERE user_id=? AND type='gider' AND date > ?", (user_id, limit_date))
    gider = cursor.fetchone()[0] or 0
    
    cursor.execute("SELECT category, SUM(amount) FROM transactions WHERE user_id=? AND type='gider' AND date > ? GROUP BY category ORDER BY SUM(amount) DESC LIMIT 1", (user_id, limit_date))
    row = cursor.fetchone()
    top_cat = f"🏆 {row[0]} (`{row[1]:.2f}` TL)" if row and row[0] else "📭 Harcama yok"
    conn.close()
    
    msg = f"📅 *Bu Hafta (Son 7 Gün)*\n\n💸 Toplam Gider: `{gider}` TL\n📂 En Çok Harcanan Kategori: {top_cat}"
    msg += f"\n\n💡 _{random.choice(FINANS_IPUCLARI)}_"
    
    await update.message.reply_text(msg, parse_mode="Markdown")

async def listele(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user_id = update.effective_user.id
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("SELECT id, type, amount, category, aciklama, date FROM transactions WHERE user_id=? ORDER BY id DESC LIMIT 5", (user_id,))
    rows = cursor.fetchall()
    conn.close()
    
    if not rows:
        await update.message.reply_text("📭 Henüz işlem yok. `/ekle` ile hemen başla! 🚀", parse_mode="Markdown")
        return
        
    msg = "📜 *Son 5 İşlem:*\n"
    for r in rows:
        icon = "💸" if r[1] == 'gider' else "💰"
        date_str = r[5].split(" ")[0] if r[5] and " " in str(r[5]) else (str(r[5]) if r[5] else "🗓️ Tarih yok")
        desc = f" - 📝 {r[4]}" if r[4] else ""
        msg += f"• `ID {r[0]}` | {icon} {r[1]}: `{r[2]}` TL ({r[3]}){desc}\n  _{date_str}_\n"
        
    msg += "\n🗑️ Silmek için: `/sil <ID>`\n🖊️ Düzenlemek için: `/duzenle <ID>`"
    await update.message.reply_text(msg, parse_mode="Markdown")

async def sil(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not context.args:
        await update.message.reply_text("🗑️ Lütfen silinecek işlemin ID'sini gir.\n📝 Örnek: `/sil 5`", parse_mode="Markdown")
        return
    try:
        item_id = int(context.args[0])
        user_id = update.effective_user.id
        conn = get_db()
        cursor = conn.cursor()
        cursor.execute("DELETE FROM transactions WHERE id=? AND user_id=?", (item_id, user_id))
        deleted = cursor.rowcount
        conn.commit()
        conn.close()
        
        if deleted > 0:
            await update.message.reply_text(f"✅ ID `{item_id}` numaralı işlem silindi. 🗑️✨", parse_mode="Markdown")
        else:
            await update.message.reply_text("❌ Bu ID'ye sahip işlem bulunamadı. 🔍")
    except ValueError:
        await update.message.reply_text("❌ Geçersiz ID. Lütfen bir sayı gir. 🔢")

async def ara_basla(update: Update, context: ContextTypes.DEFAULT_TYPE):
    keyboard = category_buttons(prefix="ara_")
    await update.message.reply_text("🔍 *Hangi kategorideki harcamaları görmek istersin?*", reply_markup=InlineKeyboardMarkup(keyboard), parse_mode="Markdown")

async def ara_kategori_secildi(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    category = query.data.replace("ara_", "")
    user_id = update.effective_user.id
    
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute(
        "SELECT id, amount, aciklama, date FROM transactions WHERE user_id=? AND type='gider' AND category=? ORDER BY date DESC LIMIT 20",
        (user_id, category)
    )
    rows = cursor.fetchall()
    conn.close()
    
    if not rows:
        await query.edit_message_text(f"📂 *{category}* kategorisinde hiç harcama bulunamadı. 📭", parse_mode="Markdown")
        return
    
    toplam = sum(r[1] for r in rows)
    msg = f"{random.choice(ARAYICI_MESAJLARI)}\n\n📂 *{category} Harcamaları (Son {len(rows)} işlem):*\n"
    for r in rows:
        date_str = r[3].split(" ")[0] if r[3] and " " in str(r[3]) else (str(r[3]) if r[3] else "🗓️ Tarih yok")
        desc = f" - 📝 {r[2]}" if r[2] else ""
        msg += f"• `ID {r[0]}` | 💸 `{r[1]:.2f}` TL{desc}\n  _{date_str}_\n"
    
    msg += f"\n💰 *Toplam:* `{toplam:.2f}` TL"
    await query.edit_message_text(msg, parse_mode="Markdown")

# --- DÜZENLEME AKIŞI ---

async def duzenle_basla(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not context.args:
        await update.message.reply_text("🖊️ Lütfen düzenlemek istediğin işlemin ID'sini gir.\n📝 Örnek: `/duzenle 5`", parse_mode="Markdown")
        return ConversationHandler.END
    
    try:
        item_id = int(context.args[0])
        user_id = update.effective_user.id
        conn = get_db()
        cursor = conn.cursor()
        cursor.execute("SELECT id, type, amount, category, aciklama FROM transactions WHERE id=? AND user_id=?", (item_id, user_id))
        row = cursor.fetchone()
        conn.close()
        
        if not row:
            await update.message.reply_text("❌ Bu ID'ye sahip işlem bulunamadı veya sana ait değil. 🔍")
            return ConversationHandler.END
            
        context.user_data['edit_id'] = item_id
        await update.message.reply_text(
            f"🖊️ *Düzenlenecek İşlem (ID {row[0]}):*\n"
            f"{row[1].upper()}: `{row[2]}` TL ({row[3]})\n"
            f"📝 Açıklama: {row[4] or '-'}\n\n"
            f"🔧 Ne değiştirmek istersin?",
            reply_markup=edit_menu_keyboard(),
            parse_mode="Markdown"
        )
        return EDIT_CHOOSE
    except ValueError:
        await update.message.reply_text("❌ Geçersiz ID. Lütfen bir sayı gir. 🔢")
        return ConversationHandler.END

async def duzenle_alan_secildi(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    action = query.data
    
    if action == "edit_bitir":
        await query.edit_message_text("✅ Düzenleme tamamlandı. Harika iş! 🎉📝")
        context.user_data.clear()
        return ConversationHandler.END
        
    if action == "edit_kategori":
        keyboard = category_buttons() + [[InlineKeyboardButton("🔙 İptal", callback_data="edit_iptal")]]
        await query.edit_message_text("📂 *Yeni kategoriyi seç:*", reply_markup=InlineKeyboardMarkup(keyboard), parse_mode="Markdown")
        return EDIT_CATEGORY
        
    if action == "edit_tutar":
        await query.edit_message_text("💵 *Yeni tutarı gir* (örn: 50 veya 50.5):" + IPTAL_HATIRLATMA, parse_mode="Markdown")
        return EDIT_AMOUNT
        
    if action == "edit_aciklama":
        await query.edit_message_text("📝 *Yeni açıklamayı yaz* (veya 'boş' yaz):" + IPTAL_HATIRLATMA, parse_mode="Markdown")
        return EDIT_DESC

async def duzenle_kategori_secildi(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    action = query.data
    
    if action == "edit_iptal":
        await query.edit_message_text("Kategori değişikliği iptal edildi. 🔙", reply_markup=edit_menu_keyboard())
        return EDIT_CHOOSE
        
    new_cat = action.replace("edit_cat_", "")
    item_id = context.user_data.get('edit_id')
    if not item_id:
        return ConversationHandler.END
        
    user_id = update.effective_user.id
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("UPDATE transactions SET category=? WHERE id=? AND user_id=?", (new_cat, item_id, user_id))
    conn.commit()
    conn.close()
    
    await query.edit_message_text(f"✅ Kategori '{new_cat}' olarak güncellendi. 📂✨", reply_markup=edit_menu_keyboard())
    return EDIT_CHOOSE

async def duzenle_tutar_alindi(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user_id = update.effective_user.id
    item_id = context.user_data.get('edit_id')
    try:
        val = float(update.message.text.replace(",", "."))
        if val <= 0:
            await update.message.reply_text("⚠️ Lütfen sıfırdan büyük bir sayı gir. 📏" + IPTAL_HATIRLATMA, parse_mode="Markdown")
            return EDIT_AMOUNT
        conn = get_db()
        cursor = conn.cursor()
        cursor.execute("UPDATE transactions SET amount=? WHERE id=? AND user_id=?", (val, item_id, user_id))
        conn.commit()
        conn.close()
        await update.message.reply_text(f"✅ Tutar `{val} TL` olarak güncellendi. 💸📝", parse_mode="Markdown")
        await update.message.reply_text(random.choice(DUZENLE_MESAJLARI))
        await update.message.reply_text("🔧 Başka bir şey değiştirmek ister misin?", reply_markup=edit_menu_keyboard())
        return EDIT_CHOOSE
    except ValueError:
        await update.message.reply_text("❌ Geçersiz tutar. Lütfen sayı gir. 🔢" + IPTAL_HATIRLATMA, parse_mode="Markdown")
        return EDIT_AMOUNT

async def duzenle_aciklama_alindi(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user_id = update.effective_user.id
    item_id = context.user_data.get('edit_id')
    desc = "" if update.message.text.lower() == "boş" else update.message.text
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("UPDATE transactions SET aciklama=? WHERE id=? AND user_id=?", (desc, item_id, user_id))
    conn.commit()
    conn.close()
    await update.message.reply_text("✅ Açıklama güncellendi. 📝🔄")
    await update.message.reply_text(random.choice(DUZENLE_MESAJLARI))
    await update.message.reply_text("🔧 Başka bir şey değiştirmek ister misin?", reply_markup=edit_menu_keyboard())
    return EDIT_CHOOSE

async def duzenle_iptal(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text("❌ İşlem düzenleme iptal edildi. 🔙")
    context.user_data.clear()
    await start(update, context)
    return ConversationHandler.END

# --- AYLIK ÖZET ---

async def aylik_ozet(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user_id = update.effective_user.id
    now = datetime.now()
    ay = now.month
    yil = now.year
    
    ay_adi = now.strftime("%B %Y")
    
    ay_baslangic = datetime(yil, ay, 1).strftime("%Y-%m-%d %H:%M:%S")
    ay_bitis = datetime(yil + 1, 1, 1).strftime("%Y-%m-%d %H:%M:%S") if ay == 12 else datetime(yil, ay + 1, 1).strftime("%Y-%m-%d %H:%M:%S")
    
    onceki_ay = 12 if ay == 1 else ay - 1
    onceki_yil = yil - 1 if ay == 1 else yil
    onceki_ay_baslangic = datetime(onceki_yil, onceki_ay, 1).strftime("%Y-%m-%d %H:%M:%S")
    onceki_ay_bitis = datetime(yil, ay, 1).strftime("%Y-%m-%d %H:%M:%S")
    
    conn = get_db()
    cursor = conn.cursor()
    
    cursor.execute("SELECT id, type, amount, category, aciklama, date FROM transactions WHERE user_id=? AND date >= ? AND date < ? ORDER BY date", (user_id, ay_baslangic, ay_bitis))
    rows = cursor.fetchall()
    
    cursor.execute("SELECT COALESCE(SUM(amount), 0) FROM transactions WHERE user_id=? AND type='gelir' AND date >= ? AND date < ?", (user_id, ay_baslangic, ay_bitis))
    toplam_gelir = cursor.fetchone()[0]
    
    cursor.execute("SELECT COALESCE(SUM(amount), 0) FROM transactions WHERE user_id=? AND type='gider' AND date >= ? AND date < ?", (user_id, ay_baslangic, ay_bitis))
    toplam_gider = cursor.fetchone()[0]
    
    cursor.execute("SELECT category, SUM(amount) FROM transactions WHERE user_id=? AND type='gider' AND date >= ? AND date < ? GROUP BY category ORDER BY SUM(amount) DESC", (user_id, ay_baslangic, ay_bitis))
    kategori_ozet = cursor.fetchall()
    
    cursor.execute("SELECT COALESCE(SUM(amount), 0) FROM transactions WHERE user_id=? AND type='gelir' AND date >= ? AND date < ?", (user_id, onceki_ay_baslangic, onceki_ay_bitis))
    onceki_gelir = cursor.fetchone()[0]
    
    cursor.execute("SELECT COALESCE(SUM(amount), 0) FROM transactions WHERE user_id=? AND type='gider' AND date >= ? AND date < ?", (user_id, onceki_ay_baslangic, onceki_ay_bitis))
    onceki_gider = cursor.fetchone()[0]
    
    cursor.execute("SELECT category, SUM(amount) FROM transactions WHERE user_id=? AND type='gider' AND date >= ? AND date < ? GROUP BY category", (user_id, onceki_ay_baslangic, onceki_ay_bitis))
    onceki_kategori = {row[0]: row[1] for row in cursor.fetchall()}
    conn.close()
    
    icgoruler = []
    if toplam_gider > 0:
        gelir_degisim = ((toplam_gelir - onceki_gelir) / onceki_gelir * 100) if onceki_gelir > 0 else 0
        gider_degisim = ((toplam_gider - onceki_gider) / onceki_gider * 100) if onceki_gider > 0 else 0
        
        if abs(gider_degisim) >= 5:
            yon = "artırdın 🔴" if gider_degisim > 0 else "azalttın 🟢"
            icgoruler.append(f"Toplam harcaman geçen aya göre %{abs(gider_degisim):.0f} {yon}")
        
        if abs(gelir_degisim) >= 5:
            yon = "arttı 📈" if gelir_degisim > 0 else "düştü 📉"
            icgoruler.append(f"Toplam gelirinin geçen aya göre %{abs(gelir_degisim):.0f} {yon}")
        
        for kat, tutar in kategori_ozet:
            onceki = onceki_kategori.get(kat, 0)
            if onceki > 0:
                degisim = ((tutar - onceki) / onceki * 100)
                if abs(degisim) >= 10:
                    if degisim > 0:
                        icgoruler.append(f"🔸 {kat} harcaman geçen aya göre %{degisim:.0f} arttı")
                    else:
                        icgoruler.append(f"🔹 {kat} harcaman geçen aya göre %{abs(degisim):.0f} azaldı")
    
    if not rows:
        await update.message.reply_text("📭 Bu ay kayıtlı işlem yok.")
        return
    
    wb = Workbook()
    ws_ozet = wb.active
    ws_ozet.title = "Aylik Ozet"
    
    baslik_font = Font(name="Calibri", bold=True, size=14, color="FFFFFF")
    baslik_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
    alt_baslik_font = Font(name="Calibri", bold=True, size=11)
    para_font = Font(name="Calibri", bold=True, size=11)
    normal_font = Font(name="Calibri", size=11)
    green_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    red_fill = PatternFill(start_color="FCE4EC", end_color="FCE4EC", fill_type="solid")
    thin_border = Border(left=Side(style="thin"), right=Side(style="thin"), top=Side(style="thin"), bottom=Side(style="thin"))
    
    ws_ozet.merge_cells("A1:B1")
    cell_baslik = ws_ozet["A1"]
    cell_baslik.value = f"AYLIK OZET - {ay_adi.upper()}"
    cell_baslik.font = baslik_font
    cell_baslik.fill = baslik_fill
    cell_baslik.alignment = Alignment(horizontal="center")
    ws_ozet["A1"].border = thin_border
    ws_ozet["B1"].border = thin_border
    
    ws_ozet["A3"] = "Toplam Gelir"
    ws_ozet["A3"].font = alt_baslik_font
    ws_ozet["B3"] = f"{toplam_gelir:,.2f} TL"
    ws_ozet["B3"].font = para_font
    ws_ozet["B3"].fill = green_fill
    ws_ozet["A3"].border = thin_border
    ws_ozet["B3"].border = thin_border
    
    ws_ozet["A4"] = "Toplam Gider"
    ws_ozet["A4"].font = alt_baslik_font
    ws_ozet["B4"] = f"{toplam_gider:,.2f} TL"
    ws_ozet["B4"].font = para_font
    ws_ozet["B4"].fill = red_fill
    ws_ozet["A4"].border = thin_border
    ws_ozet["B4"].border = thin_border
    
    net = toplam_gelir - toplam_gider
    ws_ozet["A5"] = "Net Bakiye"
    ws_ozet["A5"].font = alt_baslik_font
    ws_ozet["B5"] = f"{net:,.2f} TL"
    ws_ozet["B5"].font = para_font
    ws_ozet["B5"].fill = green_fill if net >= 0 else red_fill
    ws_ozet["A5"].border = thin_border
    ws_ozet["B5"].border = thin_border
    
    row_idx = 7
    if icgoruler:
        ws_ozet[f"A{row_idx}"] = "Aylik Karsilastirma"
        ws_ozet[f"A{row_idx}"].font = alt_baslik_font
        ws_ozet[f"A{row_idx}"].border = thin_border
        ws_ozet.merge_cells(f"A{row_idx}:B{row_idx}")
        row_idx += 1
        for ig in icgoruler:
            ws_ozet[f"A{row_idx}"] = ig
            ws_ozet[f"A{row_idx}"].font = normal_font
            ws_ozet.merge_cells(f"A{row_idx}:B{row_idx}")
            row_idx += 1
        row_idx += 1
    
    ws_ozet[f"A{row_idx}"] = "Kategori Bazli Giderler"
    ws_ozet[f"A{row_idx}"].font = alt_baslik_font
    ws_ozet[f"A{row_idx}"].border = thin_border
    ws_ozet.merge_cells(f"A{row_idx}:B{row_idx}")
    row_idx += 1
    
    ws_ozet[f"A{row_idx}"] = "Kategori"
    ws_ozet[f"A{row_idx}"].font = baslik_font
    ws_ozet[f"A{row_idx}"].fill = baslik_fill
    ws_ozet[f"B{row_idx}"] = "Tutar"
    ws_ozet[f"B{row_idx}"].font = baslik_font
    ws_ozet[f"B{row_idx}"].fill = baslik_fill
    ws_ozet[f"A{row_idx}"].border = thin_border
    ws_ozet[f"B{row_idx}"].border = thin_border
    row_idx += 1
    
    for kat, tutar in kategori_ozet:
        ws_ozet[f"A{row_idx}"] = kat
        ws_ozet[f"A{row_idx}"].font = normal_font
        ws_ozet[f"A{row_idx}"].border = thin_border
        ws_ozet[f"B{row_idx}"] = f"{tutar:,.2f} TL"
        ws_ozet[f"B{row_idx}"].font = normal_font
        ws_ozet[f"B{row_idx}"].border = thin_border
        row_idx += 1
    
    ws_ozet.column_dimensions["A"].width = 40
    ws_ozet.column_dimensions["B"].width = 20
    
    if rows:
        ws_islem = wb.create_sheet(title="Islem Listesi")
        basliklar = ["ID", "Tarih", "Tur", "Kategori", "Tutar (TL)", "Aciklama"]
        for col, baslik in enumerate(basliklar, 1):
            cell = ws_islem.cell(row=1, column=col, value=baslik)
            cell.font = baslik_font
            cell.fill = baslik_fill
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center")
        
        for r, islem in enumerate(rows, 2):
            tarih = str(islem[5]).split(" ")[0] if islem[5] else "-"
            ws_islem.cell(row=r, column=1, value=islem[0]).border = thin_border
            ws_islem.cell(row=r, column=2, value=tarih).border = thin_border
            ws_islem.cell(row=r, column=3, value=islem[1].upper()).border = thin_border
            ws_islem.cell(row=r, column=4, value=islem[3]).border = thin_border
            tutar_cell = ws_islem.cell(row=r, column=5, value=islem[2])
            tutar_cell.border = thin_border
            tutar_cell.fill = red_fill if islem[1] == "gider" else green_fill
            ws_islem.cell(row=r, column=6, value=islem[4] or "-").border = thin_border
        
        for col in range(1, 7):
            ws_islem.column_dimensions[ws_islem.cell(row=1, column=col).column_letter].width = 18
    
    dosya_adi = f"aylik_ozet_{yil}_{ay:02d}.xlsx"
    tmp_path = os.path.join(tempfile.gettempdir(), dosya_adi)
    wb.save(tmp_path)
    
    caption = f"📊 {ay_adi} aylik ozetin hazir!\n\n"
    caption += f"💰 Gelir: `{toplam_gelir:,.2f}` TL\n"
    caption += f"💸 Gider: `{toplam_gider:,.2f}` TL\n"
    caption += f"📈 Net: `{net:,.2f}` TL"
    
    if icgoruler:
        caption += "\n\n🔍 *Karsilastirma:*\n"
        for ig in icgoruler:
            caption += f"• {ig}\n"
    
    await update.message.reply_document(
        document=open(tmp_path, "rb"),
        filename=dosya_adi,
        caption=caption,
        parse_mode="Markdown"
    )
    
    os.remove(tmp_path)

# --- ANA FONKSIYON ---

def main():
    init_db()
    
    async def start_fallback(update: Update, context: ContextTypes.DEFAULT_TYPE):
        context.user_data.clear()
        await start(update, context)
        return ConversationHandler.END

    app = Application.builder().token(TOKEN).connect_timeout(20).read_timeout(20).write_timeout(20).pool_timeout(20).build()
    
    app.add_handler(CommandHandler("start", start))
    app.add_handler(CommandHandler("yardim", yardim))
    app.add_handler(CommandHandler("bakiye", bakiye))
    app.add_handler(CommandHandler("kur", kur)) 
    app.add_handler(CommandHandler("haftalik", haftalik))
    app.add_handler(CommandHandler("listele", listele))
    app.add_handler(CommandHandler("sil", sil))
    app.add_handler(CommandHandler("ara", ara_basla))
    app.add_handler(CallbackQueryHandler(ara_kategori_secildi, pattern="^ara_"))
    app.add_handler(CommandHandler("aylik_ozet", aylik_ozet))
    
    conv_handler = ConversationHandler(
        entry_points=[CommandHandler("ekle", ekle_basla)],
        states={
            WAITING_TYPE: [CallbackQueryHandler(tur_secildi)],
            WAITING_AMOUNT: [MessageHandler(filters.TEXT & ~filters.COMMAND, tutar_alindi)],
            WAITING_CATEGORY: [CallbackQueryHandler(kategori_secildi)],
            WAITING_DESC: [MessageHandler(filters.TEXT & ~filters.COMMAND, aciklama_alindi)],
        },
        fallbacks=[CommandHandler("iptal", iptal), CommandHandler("start", start_fallback)]
    )
    app.add_handler(conv_handler)
    
    conv_edit = ConversationHandler(
        entry_points=[CommandHandler("duzenle", duzenle_basla)],
        states={
            EDIT_CHOOSE: [CallbackQueryHandler(duzenle_alan_secildi, pattern="^edit_")],
            EDIT_AMOUNT: [MessageHandler(filters.TEXT & ~filters.COMMAND, duzenle_tutar_alindi)],
            EDIT_CATEGORY: [CallbackQueryHandler(duzenle_kategori_secildi, pattern="^edit_cat_|^edit_iptal")],
            EDIT_DESC: [MessageHandler(filters.TEXT & ~filters.COMMAND, duzenle_aciklama_alindi)],
        },
        fallbacks=[CommandHandler("iptal", iptal), CommandHandler("start", start_fallback)]
    )
    app.add_handler(conv_edit)

    async def error_handler(update: Update, context: ContextTypes.DEFAULT_TYPE):
        logging.error(f"Bot hatası: {context.error}")
        if update and update.message:
            await update.message.reply_text(
                "❌ Bir hata oluştu. Endişelenme, işlemi sıfırlıyorum.\n\n"
                "💡 `/ekle` veya `/duzenle` ile yeniden başlayabilirsin."
            )
            if context.user_data:
                context.user_data.clear()
    
    app.add_error_handler(error_handler)
    
    print("Bot çalışıyor...")
    app.run_polling(drop_pending_updates=True)

if __name__ == "__main__":
    main()
